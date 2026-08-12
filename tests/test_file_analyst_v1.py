from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from file_analyst import analyze_file, extract_evidence_units, result_to_bench_fragment
from intelligence_graph_db import dual_write_source_file


def _docx_bytes(paragraphs: list[str]) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "02_Cost_Breakdown"
    ws.append(["Category", "Item Description", "Vendor/Subcontractor", "Quantity", "Cost before tax", "Agency Markup", "Tax Amount", "Total including tax"])
    ws.append(["Venue & Infrastructure", "Venue Rental", "", 1, 100000, 10000, 5000, 115000])
    ws.append(["Scenic & Event Production", "Scenography", "Supplier A", 1, 200000, 20000, 10000, 230000])
    ws.append(["TOTAL", "", "", "", 300000, 30000, 15000, 345000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_briefing_file_analyst_preserves_file_and_extracts_grounded_core_claims():
    data = _docx_bytes([
        "JOB: Product Launch",
        "4. BRIEFING",
        "Quantidade de Convidados: aproximadamente 250 convidados.",
        "1. Ativação YouTube: Masterclass",
        "Adequação à Plataforma - O YouTube é o ambiente ideal para conteúdo de longa duração.",
        "Obs: gravações de conteúdo precisam ser feitas na horizontal.",
        "2. Ativação Instagram: Aesthetics & Lifestyle Gallery",
        "Adequação à Plataforma: lifestyle e conteúdo aspiracional.",
        "3. Ativação TikTok: Trend & Movement Challenge",
        "Vídeos verticais curtos e transições rápidas.",
        "FINANCEIRO: BUDGET R$1.300.000,00. Critério eliminatório.",
    ])
    result = analyze_file(
        file_name="Briefing_Produto.docx",
        data=data,
        declared_role="briefing_original",
        enable_semantic=False,
    )
    assert result.source_role == "briefing_original"
    assert len(result.evidence_units) >= 8
    claims = {(c.subject_key, c.predicate): c for c in result.claims}
    assert claims[("project", "budget_max")].value_numeric == 1_300_000.0
    assert claims[("project", "expected_attendees")].value_numeric == 250.0
    assert claims[("youtube_requirement", "required_platform_behavior")].value_text == "horizontal"
    assert claims[("instagram_requirement", "required_platform_behavior")].value_text == "lifestyle / aspirational content"
    assert claims[("tiktok_requirement", "required_platform_behavior")].value_text == "vertical"
    assert all(c.evidence_refs for c in result.claims)


def test_cost_sheet_creates_line_entities_and_proposed_not_actual_total():
    data = _xlsx_bytes()
    result = analyze_file(
        file_name="Quotation.xlsx",
        data=data,
        declared_role="detailed_costs",
        enable_semantic=False,
    )
    assert result.source_role == "detailed_costs"
    proposed = [c for c in result.claims if c.predicate == "proposed_total"]
    actual = [c for c in result.claims if c.predicate == "actual_total"]
    assert proposed
    assert round(float(proposed[0].value_numeric or 0), 2) == 345000.00
    assert actual == []
    assert any(e.entity_type == "financial_line_item" for e in result.entities)
    assert any(u.unit_type == "row" for u in result.evidence_units)


def test_image_without_semantic_keeps_evidence_without_inventing_transcription():
    data = b"\xff\xd8\xff\xe0" + b"fake-jpeg-bytes"
    result = analyze_file(
        file_name="feedback.jpeg",
        data=data,
        mime_type="image/jpeg",
        declared_role="feedback_approval",
        enable_semantic=False,
    )
    assert result.source_role == "feedback_approval"
    assert len(result.evidence_units) == 1
    assert result.evidence_units[0].unit_type == "image"
    assert result.claims == []
    assert result.metadata == {}


def test_bench_fragment_prefixes_provenance_refs():
    data = _docx_bytes(["BUDGET R$100.000,00", "100 convidados"])
    result = analyze_file(
        file_name="Briefing.docx",
        data=data,
        declared_role="briefing_original",
        enable_semantic=False,
    )
    fragment = result_to_bench_fragment(result, "briefing")
    assert fragment["source_role"] == "briefing_original"
    assert all(
        ref.startswith("briefing:")
        for claim in fragment["claims"]
        for ref in claim.get("evidence_refs", [])
    )


class _MissingFoundationClient:
    def table(self, _name):
        raise RuntimeError("relation does not exist")


def test_dual_write_is_fail_open_when_foundation_is_not_installed():
    result = dual_write_source_file(
        _MissingFoundationClient(),
        {"id": "sf1", "project_id": "p1", "file_name": "x.txt", "mime_type": "text/plain"},
        source_bytes=b"hello",
        enable_semantic=False,
    )
    assert result["status"] == "skipped_foundation_missing"
